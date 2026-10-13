from datetime import date, time, timedelta

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from .models import (
    Cabinet, Child, DuomenuImportas, Klase, LeistinasEmail, Profile, Reservation, Subject, Teacher, Vadovas,
    VadovoLaikas, VadovoRezervacija, WorkingHours,
)
from .importas import importuoti_mokytojus
from .duomenu_importas import nuskaityti_mokytoju_sarasa, perziureti_importa, vykdyti_importa
from .views import sudaryti_laikus


class VadovoIrTevuSusikirtimai(TestCase):
    def setUp(self):
        self.diena = date.today() + timedelta(days=7)
        self.kab = Cabinet.objects.create(pavadinimas="101")

        self.mokytojas = Teacher.objects.create(vardas="Ona", pavarde="Onaitė", email="ona@test.lt")
        self.wh = WorkingHours.objects.create(
            date=self.diena, start_time=time(15, 0), end_time=time(16, 0),
            interval=10, tipas="individualus",
        )
        self.wh.mokytojai.add(self.mokytojas)

        self.vadovas = Vadovas.objects.create(vardas="Jonas", pavarde="Jonaitis", email="jonas@test.lt")
        self.vl = VadovoLaikas.objects.create(
            vadovas=self.vadovas, date=self.diena, start_time=time(15, 0),
            end_time=time(16, 0), interval=15, cabinet=self.kab,
        )

        self.tevas = User.objects.create_user("tevas", email="tevas@test.lt")
        vaikas = Child.objects.create(user=self.tevas, first_name="Petras", last_name="P")
        Profile.objects.create(user=self.tevas, active_child=vaikas)

    def rezervuoti_tevui(self, laikas):
        self.client.force_login(self.tevas)
        return self.client.post(reverse("reserve_timeslot", args=[self.mokytojas.id]), {
            "date": self.diena.isoformat(), "times": [laikas],
            "tipas": "individualus", "confirm": "yes",
        })

    def test_tevai_negali_rezervuoti_kai_mokytojas_pas_vadova(self):
        # pokalbis su vadovu 15:00-15:15 uzima tevu laikus 15:00 ir 15:10
        VadovoRezervacija.objects.create(
            vadovas=self.vadovas, mokytojas=self.mokytojas, date=self.diena, time=time(15, 0),
        )
        self.rezervuoti_tevui("15:10")
        self.assertFalse(Reservation.objects.exists())

        self.rezervuoti_tevui("15:20")
        self.assertTrue(Reservation.objects.filter(time=time(15, 20)).exists())

    def test_grafike_laikas_rodomas_uzimtu(self):
        VadovoRezervacija.objects.create(
            vadovas=self.vadovas, mokytojas=self.mokytojas, date=self.diena, time=time(15, 0),
        )
        laikai = {s["time"]: s["status"] for s in sudaryti_laikus(self.mokytojas, self.diena, [self.wh], show_names=False)}
        self.assertEqual(laikai["15:00"], "busy")
        self.assertEqual(laikai["15:10"], "busy")
        self.assertEqual(laikai["15:20"], "free")

    def test_mokytojas_negali_pas_vadova_kai_uzsiregistrave_tevai(self):
        Reservation.objects.create(
            user=self.tevas, teacher=self.mokytojas, date=self.diena, time=time(15, 10),
        )
        mokytojo_user = User.objects.create_user("ona", email="ona@test.lt")
        self.client.force_login(mokytojo_user)
        self.client.post(reverse("pokalbiai_vadovas_rezervuoti", args=[self.vadovas.id]), {
            "date": self.diena.isoformat(), "time": "15:00", "confirm": "yes",
        })
        self.assertFalse(VadovoRezervacija.objects.exists())


class BeIntervalo(TestCase):
    def setUp(self):
        self.diena = date.today() + timedelta(days=7)
        self.kab = Cabinet.objects.create(pavadinimas="101")
        self.mokytojas = Teacher.objects.create(vardas="Ona", pavarde="Onaitė", email="ona@test.lt")
        self.vadovas = Vadovas.objects.create(vardas="Jonas", pavarde="Jonaitis", email="jonas@test.lt")

    def test_numatytasis_be_intervalo(self):
        self.client.force_login(User.objects.create_user("ona", email="ona@test.lt"))
        self.client.post(reverse("teacher_add_workinghours"), {
            "date": self.diena.isoformat(), "start_time": "15:00", "end_time": "16:30",
            "interval": "0", "break_start": "15:30", "break_end": "15:45",
        })
        wh = WorkingHours.objects.get()
        self.assertEqual(wh.interval, 0)
        self.assertEqual(wh.trukme, 90)
        # pertrauka is formos neissaugoma
        self.assertFalse(wh.breaks.exists())

        laikai = sudaryti_laikus(self.mokytojas, self.diena, [wh])
        self.assertEqual([s["time"] for s in laikai], ["15:00"])

        self.client.post(reverse("teacher_add_break", args=[wh.id]), {
            "start_time": "15:30", "end_time": "15:45",
        })
        self.assertFalse(wh.breaks.exists())

    def test_vadovas_be_intervalo(self):
        self.client.force_login(User.objects.create_user("jonas", email="jonas@test.lt"))
        self.client.post(reverse("vadovas_add_laikas"), {
            "date": self.diena.isoformat(), "start_time": "15:00", "end_time": "16:00",
            "interval": "0", "cabinet": self.kab.id,
        })
        vl = VadovoLaikas.objects.get()
        self.assertEqual(vl.trukme, 60)

        self.client.post(reverse("vadovas_add_break", args=[vl.id]), {
            "start_time": "15:30", "end_time": "15:45",
        })
        self.assertFalse(vl.breaks.exists())

        # mokytojas gali uzsiregistruoti tik 15:00, o tevai tuo metu (visa valanda) - ne
        wh = WorkingHours.objects.create(
            date=self.diena, start_time=time(15, 0), end_time=time(16, 0), interval=15,
        )
        wh.mokytojai.add(self.mokytojas)
        self.client.force_login(User.objects.create_user("ona", email="ona@test.lt"))
        url = reverse("pokalbiai_vadovas_rezervuoti", args=[self.vadovas.id])
        self.client.post(url, {"date": self.diena.isoformat(), "time": "15:15", "confirm": "yes"})
        self.assertFalse(VadovoRezervacija.objects.exists())
        self.client.post(url, {"date": self.diena.isoformat(), "time": "15:00", "confirm": "yes"})
        self.assertTrue(VadovoRezervacija.objects.exists())

        laikai = {s["time"]: s["status"] for s in sudaryti_laikus(self.mokytojas, self.diena, [wh], show_names=False)}
        self.assertEqual(laikai["15:45"], "busy")


class MokytojuImportas(TestCase):
    def _xlsx(self, eilutes):
        import io
        import openpyxl
        knyga = openpyxl.Workbook()
        for eil in eilutes:
            knyga.active.append(eil)
        buf = io.BytesIO()
        knyga.save(buf)
        buf.seek(0)
        return buf

    def test_antraste_el_pasto_adresas_ir_nr_stulpelis(self):
        # tokia forma kaip mokyklos G10 faile: nr. kaip float, papildomi stulpeliai
        failas = self._xlsx([
            ["nr.", "Vardas", "Pavardė", "El. pašto adresas", "Dalykas", "Kabinetas"],
            [1.0, "Agnė", "Laskytė", "agne.laskyte@herojus.lt", "Psichologė", "3a"],
            [2.0, "Asta", "Budrikaitė", "Asta.Budrikaite@herojus.lt"],
            [None],
        ])
        ataskaita = importuoti_mokytojus(failas, "G10.xlsx")

        self.assertEqual(ataskaita["klaidos"], [])
        self.assertEqual(ataskaita["sukurta"], 2)
        self.assertEqual(ataskaita["praleista"], 0)
        self.assertTrue(Teacher.objects.filter(
            vardas="Asta", pavarde="Budrikaitė", email="asta.budrikaite@herojus.lt",
        ).exists())

    def test_be_antrasciu_su_nr(self):
        failas = self._xlsx([[1.0, "Ona", "Onaitė", "ona@test.lt"]])
        ataskaita = importuoti_mokytojus(failas, "m.xlsx")
        self.assertEqual(ataskaita["sukurta"], 1)

    def test_mokytoju_skiltyje_ikelimas_kaip_anksciau(self):
        self.client.force_login(User.objects.create_superuser("admin", "admin@test.lt", "x"))
        failas = self._xlsx([["Vardas", "Pavardė", "El. paštas"], ["Jonas", "Jonaitis", "jonas@herojus.lt"]])
        failas.name = "m.xlsx"
        r = self.client.post(reverse("admin:reservations_teacher_importuoti"), {"failas": failas})
        self.assertRedirects(r, reverse("admin:reservations_teacher_changelist"))
        self.assertTrue(Teacher.objects.filter(email="jonas@herojus.lt").exists())


def xlsx_failas(*lapai):
    """lapai: (pavadinimas, eilutes) poros"""
    import io
    import openpyxl
    knyga = openpyxl.Workbook()
    knyga.remove(knyga.active)
    for pavadinimas, eilutes in lapai:
        lapas = knyga.create_sheet(pavadinimas)
        for eil in eilutes:
            lapas.append(eil)
    buf = io.BytesIO()
    knyga.save(buf)
    buf.seek(0)
    buf.name = "mokytojai.xlsx"
    return buf


ANTRASTE = ["nr.", "Vardas", "Pavardė", "El. paštas", "Dalykas", "Kabinetas", "Auklėjama klasė"]


def importuoti(eilutes, antraste=ANTRASTE):
    """Visas kelias kaip admine: nuskaityti, sudaryti plana, ivykdyti"""
    nuskaityta = nuskaityti_mokytoju_sarasa(xlsx_failas(("Mokytojai", [antraste] + eilutes)), "m.xlsx")
    if nuskaityta["klaidos"]:
        return nuskaityta
    planas = perziureti_importa(nuskaityta["eilutes"])
    return vykdyti_importa(nuskaityta["eilutes"], planas["parasas"])


def db_busena():
    return (
        list(Teacher.objects.order_by("id").values_list("id", "vardas", "pavarde", "email", "kabinetas_id")),
        list(Teacher.dalykai.through.objects.order_by("id").values_list("teacher_id", "subject_id")),
        list(Klase.objects.order_by("id").values_list("id", "pavadinimas", "aukletojas_id")),
        Cabinet.objects.count(), Subject.objects.count(),
    )


class MokytojuSarasoImportas(TestCase):
    def setUp(self):
        self.kab = Cabinet.objects.create(pavadinimas="Mundus klasė")
        self.matematika = Subject.objects.create(pavadinimas="Matematika")
        self.k5b = Klase.objects.create(pavadinimas="5B")
        self.k6a = Klase.objects.create(pavadinimas="6A")

        # esamas, atpazistamas pagal el. pasta (didziosios raides nesvarbu)
        self.brigita = Teacher.objects.create(vardas="Brigita", pavarde="Misevičiutė", email="Brigita.Miseviciute@herojus.lt")
        self.brigita.dalykai.add(self.matematika)
        # esamas be el. pasto - atpazistamas pagal varda ir pavarde
        self.dainora = Teacher.objects.create(vardas="Dainora", pavarde="Stašaitienė", kabinetas=self.kab)
        # jau turi 6A - faile 6A atiduodama kitam
        self.senas = Teacher.objects.create(vardas="Petras", pavarde="Senas", email="petras@herojus.lt")
        self.k6a.aukletojas = self.senas
        self.k6a.save()
        # nepasikeite
        self.ona = Teacher.objects.create(vardas="Ona", pavarde="Onaitė", email="ona@herojus.lt")

    def test_pilnas_importas(self):
        planas = importuoti([
            [1, "Agnė", "Laskytė", "agne.laskyte@herojus.lt", "Psichologė", "HAU klasė 3a.", None],
            [2, "Brigita", "Misevičiūtė", "brigita.miseviciute@herojus.lt", "Prancūzų kalbos mokytoja, klasės vadovė", "mundus  klasė", "5B"],
            [3, "Dainora", "Stašaitienė", "dainora.stasaitiene@herojus.lt", "Matematikos mokytoja, klasės vadovė", None, "6A"],
            [4, "Augustina", "Mackelaitė", "augustina.mackelaite@herojus.lt", "Klasės vadovė", "Augustinos klasė", "1 Augustinos | Žurnalistai"],
            [None, None, None, None, None, None, None],
            [5, "Ona", "Onaitė", "ona@herojus.lt", None, None, None],
        ])
        self.assertTrue(planas["ivykdyta"], planas.get("klaidos"))
        self.assertEqual(len(planas["nauji"]), 2)
        self.assertEqual(len(planas["atnaujinami"]), 2)  # Brigita, Dainora
        self.assertEqual(planas["nepakite"], 1)
        # Petro faile nera, bet kad jis netenka 6A - matosi klasiu sarase
        self.assertIn(
            {"klase": "6A", "buvo": "Petras Senas", "bus": "Dainora Stašaitienė", "nauja": False},
            planas["klasiu_pokyciai"],
        )

        agne = Teacher.objects.get(email="agne.laskyte@herojus.lt")
        self.assertEqual(agne.kabinetas.pavadinimas, "HAU klasė 3a.")
        self.assertEqual([d.pavadinimas for d in agne.dalykai.all()], ["Psichologė"])

        # pavarde pataisyta, dalykai pakeisti, esamas kabinetas atpazintas nepaisant tarpu/raidziu
        self.brigita.refresh_from_db()
        self.assertEqual(self.brigita.pavarde, "Misevičiūtė")
        self.assertEqual(self.brigita.email, "Brigita.Miseviciute@herojus.lt")
        self.assertEqual(sorted(d.pavadinimas for d in self.brigita.dalykai.all()), ["Klasės vadovė", "Prancūzų kalbos mokytoja"])
        self.assertEqual(self.brigita.kabinetas, self.kab)
        self.assertEqual(Cabinet.objects.filter(pavadinimas__icontains="mundus").count(), 1)

        # be el. pasto - rastas pagal varda, el. pastas pridetas, tuscias kabinetas nepakeistas
        self.dainora.refresh_from_db()
        self.assertEqual(self.dainora.email, "dainora.stasaitiene@herojus.lt")
        self.assertEqual(self.dainora.kabinetas, self.kab)
        self.assertEqual(Teacher.objects.filter(vardas="Dainora").count(), 1)

        # klases
        self.k5b.refresh_from_db()
        self.k6a.refresh_from_db()
        self.assertEqual(self.k5b.aukletojas, self.brigita)
        self.assertEqual(self.k6a.aukletojas, self.dainora)
        nauja = Klase.objects.get(pavadinimas="1 Augustinos | Žurnalistai")
        self.assertEqual(nauja.aukletojas.email, "augustina.mackelaite@herojus.lt")
        # "Klasės vadovė" sukurtas viena karta, nors faile tris kartus
        self.assertEqual(Subject.objects.filter(pavadinimas="Klasės vadovė").count(), 1)

        # kartojant ta pati faila - nieko nebekeiciama
        busena = db_busena()
        planas = importuoti([
            [1, "Agnė", "Laskytė", "agne.laskyte@herojus.lt", "Psichologė", "HAU klasė 3a.", None],
            [2, "Brigita", "Misevičiūtė", "brigita.miseviciute@herojus.lt", "Prancūzų kalbos mokytoja, klasės vadovė", "mundus  klasė", "5B"],
        ])
        self.assertEqual((len(planas["nauji"]), len(planas["atnaujinami"]), planas["nepakite"]), (0, 0, 2))
        self.assertEqual(db_busena(), busena)

    def test_klases_perdavimas_matosi_plane(self):
        nuskaityta = nuskaityti_mokytoju_sarasa(xlsx_failas(("L", [ANTRASTE, [1, "Ona", "Onaitė", "ona@herojus.lt", None, None, "6a"]])), "m.xlsx")
        planas = perziureti_importa(nuskaityta["eilutes"])
        self.assertEqual(planas["klasiu_pokyciai"], [{"klase": "6A", "buvo": "Petras Senas", "bus": "Ona Onaitė", "nauja": False}])
        self.assertEqual(planas["atnaujinami"][0]["pakeitimai"], [("Auklėjama klasė", "—", "6A")])
        self.assertEqual(planas["naujos_klases"], [])
        # planas nieko nekeicia
        self.k6a.refresh_from_db()
        self.assertEqual(self.k6a.aukletojas, self.senas)

    def test_antraste_be_nr_stulpelio_ir_csv(self):
        import io
        csv = io.BytesIO(
            "Vardas;Pavardė;El. paštas;Dalykas;Kabinetas;Auklėjama klasė\n"
            "Jonas;Jonaitis;jonas@herojus.lt;Fizika;;\n".encode("utf-8-sig")
        )
        nuskaityta = nuskaityti_mokytoju_sarasa(csv, "m.csv")
        self.assertEqual(nuskaityta["klaidos"], [])
        self.assertEqual(nuskaityta["eilutes"][0]["dalykai"], ["Fizika"])

    def test_netinkama_forma_nieko_nekeicia(self):
        busena = db_busena()
        eilute = [[1, "Agnė", "Laskytė", "agne@herojus.lt", "Psichologė", "", ""]]
        blogos = [
            ["nr.", "Vardas", "Pavardė", "El. paštas", "Dalykas", "Kabinetas"],  # truksta stulpelio
            ["nr.", "Pavardė", "Vardas", "El. paštas", "Dalykas", "Kabinetas", "Auklėjama klasė"],  # kita tvarka
            ["nr.", "Vardas", "Pavardė", "El. paštas", "Dalykas", "Kabinetas", "Auklėjama klasė", "Pastabos"],  # papildomas
            ["nr.", "Vardas", "Pavardė", "El. pašto adresas", "Dalykas", "Kabinetas", "Auklėjama klasė"],  # kitas pavadinimas
            ["x", "y", "Vardas", "Pavardė", "El. paštas", "Dalykas", "Kabinetas", "Auklėjama klasė"],  # per daug priekyje
        ]
        for antraste in blogos:
            rezultatas = importuoti(eilute, antraste)
            self.assertTrue(rezultatas["klaidos"], antraste)
            self.assertIn("Netinkama failo forma", rezultatas["klaidos"][0])
        # pavadinimo eilute virs lenteles
        rezultatas = nuskaityti_mokytoju_sarasa(xlsx_failas(("L", [["Mokytojų sąrašas"], ANTRASTE] + eilute)), "m.xlsx")
        self.assertTrue(rezultatas["klaidos"])
        self.assertEqual(db_busena(), busena)

    def test_antrastes_tarpai_ir_raides_nesvarbu(self):
        antraste = ["VARDAS ", "Pavardė", "El.paštas", "dalykas", "Kabinetas", "Auklėjama  klasė"]
        rezultatas = importuoti([["Agnė", "Laskytė", "agne@herojus.lt", "", "", ""]], antraste)
        self.assertTrue(rezultatas["ivykdyta"])

    def test_klaida_eiluteje_nieko_nekeicia(self):
        busena = db_busena()
        atvejai = [
            [[1, "Agnė", "Laskytė", "agne@herojus", "", "", ""]],  # netinkamas el. pastas
            [[1, "Agnė", "", "agne@herojus.lt", "", "", ""]],  # nera pavardes
            [[1, "Agnė", "Laskytė", "", "", "", ""]],  # nera el. pasto
            [[1, "Agnė", "Laskytė", "agne@herojus.lt", "", "", ""], [2, "Agnė", "Kita", "AGNE@herojus.lt", "", "", ""]],
            [[1, "Agnė", "Laskytė", "agne@herojus.lt", "", "", "5B"], [2, "Ona", "Onaitė", "ona@herojus.lt", "", "", "5b"]],
            [[1, "Agnė", "Laskytė", "agne@herojus.lt", "", "", "K" * 51]],  # per ilgas klases pavadinimas
            # gera eilute + bloga - neirasoma ir gera
            [[1, "Jonas", "Jonaitis", "jonas@herojus.lt", "", "", ""], [2, "Agnė", "Laskytė", "agne@", "", "", ""]],
        ]
        for eilutes in atvejai:
            rezultatas = importuoti(eilutes)
            self.assertTrue(rezultatas["klaidos"], eilutes)
            self.assertFalse(rezultatas.get("ivykdyta"))
        self.assertEqual(db_busena(), busena)

    def test_dublikatas_sujungiamas_ir_importuojamas(self):
        dublikatas = Teacher.objects.create(vardas="Ona", pavarde="Dublikatė", email="ONA@herojus.lt")
        rezultatas = importuoti([[1, "Ona", "Onaitė", "ona@herojus.lt", "Fizika", "", ""]])
        self.assertTrue(rezultatas["ivykdyta"], rezultatas["klaidos"])
        self.assertEqual(len(rezultatas["dublikatai"]), 1)
        self.assertFalse(Teacher.objects.filter(pk=dublikatas.pk).exists())
        self.ona.refresh_from_db()
        self.assertEqual([d.pavadinimas for d in self.ona.dalykai.all()], ["Fizika"])

    def test_dvi_eilutes_tam_paciam_mokytojui(self):
        # Dainora sistemoje be el. pasto - abi eilutes ja atpazintu pagal varda
        busena = db_busena()
        rezultatas = importuoti([
            [1, "Dainora", "Stašaitienė", "dainora1@herojus.lt", "", "", ""],
            [2, "Dainora", "Stašaitienė", "dainora2@herojus.lt", "", "", ""],
        ])
        self.assertIn("jau atnaujinamas pagal 2 eilutę", rezultatas["klaidos"][0])
        self.assertEqual(db_busena(), busena)

    def test_bendravardis_su_kitu_el_pastu_ispejimas(self):
        nuskaityta = nuskaityti_mokytoju_sarasa(xlsx_failas(("L", [ANTRASTE, [1, "Ona", "Onaitė", "kita.ona@herojus.lt", "", "", ""]])), "m.xlsx")
        planas = perziureti_importa(nuskaityta["eilutes"])
        self.assertEqual(len(planas["nauji"]), 1)
        self.assertIn("ona@herojus.lt", planas["ispejimai"][0])

    def test_keli_lapai(self):
        tinkamas = [ANTRASTE, [1, "Jonas", "Jonaitis", "jonas@herojus.lt", "", "", ""]]
        nuskaityta = nuskaityti_mokytoju_sarasa(xlsx_failas(("Kita", [["Kažkas"]]), ("Mokytojai", tinkamas)), "m.xlsx")
        self.assertEqual(nuskaityta["klaidos"], [])
        self.assertEqual(nuskaityta["lapas"], "Mokytojai")

        nuskaityta = nuskaityti_mokytoju_sarasa(xlsx_failas(("A", tinkamas), ("B", tinkamas)), "m.xlsx")
        self.assertIn("keli lapai", nuskaityta["klaidos"][0])

    def test_ne_excel_failas(self):
        import io
        rezultatas = nuskaityti_mokytoju_sarasa(io.BytesIO(b"ne excel"), "m.xlsx")
        self.assertIn("nepavyko perskaityti", rezultatas["klaidos"][0])

    def test_pasikeitus_db_po_perziuros_neirasoma(self):
        nuskaityta = nuskaityti_mokytoju_sarasa(xlsx_failas(("L", [ANTRASTE, [1, "Ona", "Onaitė", "ona@herojus.lt", "", "", "5B"]])), "m.xlsx")
        planas = perziureti_importa(nuskaityta["eilutes"])
        # kol adminas ziurejo perziura, kitas adminas priskyre 5B
        self.k5b.aukletojas = self.senas
        self.k5b.save()
        rezultatas = vykdyti_importa(nuskaityta["eilutes"], planas["parasas"])
        self.assertFalse(rezultatas["ivykdyta"])
        self.k5b.refresh_from_db()
        self.assertEqual(self.k5b.aukletojas, self.senas)


class MokytojuImportasAdmine(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("admin", "admin@test.lt", "x")
        self.client.force_login(self.admin)
        self.url = reverse("admin:reservations_duomenuimportas_changelist")
        self.patvirtinti = reverse("admin:reservations_duomenuimportas_patvirtinti")

    def _ikelti(self, eilutes, antraste=ANTRASTE):
        return self.client.post(self.url, {"failas": xlsx_failas(("L", [antraste] + eilutes))})

    def test_perziura_ir_patvirtinimas(self):
        r = self._ikelti([[1, "Agnė", "Laskytė", "agne@herojus.lt", "Psichologė", "HAU klasė 3a.", "5B"]])
        self.assertContains(r, "Nauji mokytojai (1)")
        self.assertContains(r, "jos bus sukurtos")
        # perziura dar nieko nekeicia
        self.assertFalse(Teacher.objects.exists())
        self.assertFalse(Klase.objects.exists())

        r = self.client.post(self.patvirtinti)
        self.assertRedirects(r, self.url)
        mokytojas = Teacher.objects.get()
        self.assertEqual(Klase.objects.get().aukletojas, mokytojas)

        # antra karta patvirtinti nebegalima - duomenys istrinti is sesijos
        r = self.client.post(self.patvirtinti, follow=True)
        self.assertContains(r, "Niekas nepakeista")
        self.assertEqual(Teacher.objects.count(), 1)

    def test_bloga_forma_rodo_klaida(self):
        r = self._ikelti([["Agnė", "Laskytė", "agne@herojus.lt"]], ["Vardas", "Pavardė", "El. paštas"])
        self.assertContains(r, "niekas nepakeista")
        self.assertContains(r, "Netinkama failo forma")
        self.assertFalse(Teacher.objects.exists())
        self.assertNotIn("duomenu_importas", self.client.session)

    def test_be_teisiu_negalima(self):
        darbuotojas = User.objects.create_user("darb", "d@test.lt", "x", is_staff=True)
        self.client.force_login(darbuotojas)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.client.post(self.patvirtinti).status_code, 403)

    def test_skiltis_rezervaciju_sarase(self):
        r = self.client.get(reverse("admin:app_list", args=["reservations"]))
        self.assertContains(r, "Importuoti duomenys")
        self.assertContains(r, f'href="{self.url}"')
        # puslapis - ikelimo forma, ne sarasas
        r = self.client.get(self.url)
        self.assertContains(r, "Pavardė")
        self.assertContains(r, "Patikrinti dublikatus")

    def test_tik_dublikatai(self):
        k1 = Klase.objects.create(pavadinimas="6A")
        k2 = Klase.objects.create(pavadinimas="6a")
        r = self.client.post(self.url, {"tik_dublikatai": "1"})
        self.assertContains(r, "Dublikatų patikrinimas")
        self.assertContains(r, "Dublikatai – bus sujungti")
        self.assertEqual(Klase.objects.count(), 2)
        r = self.client.post(self.patvirtinti)
        self.assertRedirects(r, self.url)
        self.assertEqual(list(Klase.objects.values_list("pk", flat=True)), [k1.pk])

        # dublikatu nebera
        r = self.client.post(self.url, {"tik_dublikatai": "1"})
        self.assertContains(r, "Dublikatų nerasta")


class LietuviskasRikiavimas(TestCase):
    def test_mokytojai_ir_klases(self):
        for pavarde in ["Žukauskas", "Zigmas", "Černiauskienė", "Cibulskis", "Dagilė", "Šlubelytė", "Sabaliauskas", "Tamulis"]:
            Teacher.objects.create(vardas="A", pavarde=pavarde)
        self.assertEqual(
            [m.pavarde for m in Teacher.objects.all()],
            ["Cibulskis", "Černiauskienė", "Dagilė", "Sabaliauskas", "Šlubelytė", "Tamulis", "Zigmas", "Žukauskas"],
        )
        for pav in ["10S", "6f", "11V", "9D", "6A", "1 Augustinos | Žurnalistai"]:
            Klase.objects.create(pavadinimas=pav)
        self.assertEqual(
            [k.pavadinimas for k in Klase.objects.all()],
            ["1 Augustinos | Žurnalistai", "6A", "6f", "9D", "10S", "11V"],
        )


class AdminSarasuRikiavimas(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_superuser("admin", "admin@test.lt", "x"))
        diena = date.today() + timedelta(days=7)
        kab = Cabinet.objects.create(pavadinimas="Šviesa")
        k10 = Klase.objects.create(pavadinimas="10S")
        k6 = Klase.objects.create(pavadinimas="6F")
        fizika = Subject.objects.create(pavadinimas="Fizika")
        self.z = Teacher.objects.create(vardas="Ona", pavarde="Žukauskaitė", email="z@test.lt", kabinetas=kab)
        self.c = Teacher.objects.create(vardas="Ona", pavarde="Čepaitė", email="c@test.lt")
        self.b = Teacher.objects.create(vardas="Ona", pavarde="Bitė", email="b@test.lt")
        self.z.klases.add(k6, k10)
        self.c.klases.add(k10)
        self.z.dalykai.add(fizika)
        k6.aukletojas = self.c
        k6.save()
        tevai = User.objects.create_user("tevai", "t@test.lt")
        vaikas = Child.objects.create(user=tevai, first_name="Jonas", last_name="Jonaitis", klase=k6)
        Profile.objects.create(user=tevai, active_child=vaikas)
        wh = WorkingHours.objects.create(date=diena, start_time=time(15, 0), end_time=time(16, 0), interval=15)
        wh.mokytojai.add(self.z, self.c)
        Reservation.objects.create(user=tevai, teacher=self.z, date=diena, time=time(15, 0), child=vaikas, reserved_class="6F")
        vadovas = Vadovas.objects.create(vardas="Jonas", pavarde="Vadovas", email="v@test.lt")
        VadovoLaikas.objects.create(vadovas=vadovas, date=diena, start_time=time(15, 0), end_time=time(16, 0), cabinet=kab)
        VadovoRezervacija.objects.create(vadovas=vadovas, mokytojas=self.z, date=diena, time=time(15, 0), cabinet=kab)
        LeistinasEmail.objects.create(email="tevai@test.lt", pastaba="Jonas")

    def test_visi_stulpeliai_visuose_sarasuose(self):
        from django.contrib import admin as django_admin
        for modelis, modelio_admin in django_admin.site._registry.items():
            # importo skiltis - ne sarasas
            if modelis._meta.app_label != "reservations" or modelis is DuomenuImportas:
                continue
            url = reverse(f"admin:reservations_{modelis._meta.model_name}_changelist")
            for i in range(len(modelio_admin.list_display)):
                for o in (str(i), f"-{i}"):
                    r = self.client.get(url, {"o": o})
                    self.assertEqual(r.status_code, 200, f"{url}?o={o}")
                    cl = r.context["cl"]
                    self.assertTrue(cl.result_list, url)
                    # kiekvienas stulpelis turi buti rikiuojamas
                    for laukas in cl.list_display:
                        if laukas != "action_checkbox":
                            self.assertIsNotNone(cl.get_ordering_field(laukas), f"{url} {laukas}")

    def _eiles(self, url, o, **kitas):
        r = self.client.get(url, {"o": o, **kitas})
        return [obj.pk for obj in r.context["cl"].result_list]

    def test_mokytojai_pagal_pavarde_ir_klases(self):
        url = reverse("admin:reservations_teacher_changelist")
        # 2 = pavarde, 4 = klases
        self.assertEqual(self._eiles(url, "2"), [self.b.pk, self.c.pk, self.z.pk])
        self.assertEqual(self._eiles(url, "-2"), [self.z.pk, self.c.pk, self.b.pk])
        # pagal pirma (maziausia) klase: Z turi 6F, C - 10S, B klasiu neturi - gale
        self.assertEqual(self._eiles(url, "4"), [self.z.pk, self.c.pk, self.b.pk])
        # su filtru pagal klase eilutes nesidubliuoja
        k10 = Klase.objects.get(pavadinimas="10S")
        self.assertEqual(self._eiles(url, "4", klases__id__exact=k10.pk), [self.z.pk, self.c.pk])

    def test_darbo_laikai_pagal_mokytojus(self):
        url = reverse("admin:reservations_workinghours_changelist")
        diena = date.today() + timedelta(days=8)
        wh_b = WorkingHours.objects.create(date=diena, start_time=time(15, 0), end_time=time(16, 0), interval=15)
        wh_b.mokytojai.add(self.b)
        wh_zc = WorkingHours.objects.exclude(pk=wh_b.pk).get()
        # 6 = mokytojai: Bitė < Čepaitė (pirma bloke su Žukauskaite)
        self.assertEqual(self._eiles(url, "6"), [wh_b.pk, wh_zc.pk])
        self.assertEqual(self._eiles(url, "-6"), [wh_zc.pk, wh_b.pk])

    def test_mokytojo_trynimas_is_saraso_veikia(self):
        url = reverse("admin:reservations_teacher_changelist")
        r = self.client.post(url, {"action": "delete_selected", "_selected_action": [self.b.pk], "post": "yes"})
        self.assertEqual(r.status_code, 302)
        self.assertFalse(Teacher.objects.filter(pk=self.b.pk).exists())
        self.assertEqual(Teacher.objects.count(), 2)


class TevuMokytojuSarasas(TestCase):
    def test_dalykininku_sarasas_lietuviska_tvarka(self):
        diena = date.today() + timedelta(days=7)
        k6 = Klase.objects.create(pavadinimas="6F")
        wh = WorkingHours.objects.create(
            date=diena, start_time=time(15, 0), end_time=time(16, 0), interval=15, tipas="dalykininku",
        )
        for pavarde in ["Žukauskas", "Černiauskienė", "Dagilė", "Cibulskis"]:
            wh.mokytojai.add(Teacher.objects.create(vardas="A", pavarde=pavarde))
        tevai = User.objects.create_user("tevai", "t@test.lt")
        vaikas = Child.objects.create(user=tevai, first_name="Jonas", last_name="Jonaitis", klase=k6)
        Profile.objects.create(user=tevai, active_child=vaikas)
        self.client.force_login(tevai)
        r = self.client.get(reverse("teacher_schedule", args=[diena.isoformat()]), {"tipas": "dalykininku"})
        self.assertEqual(
            [d["teacher"].pavarde for d in r.context["data"]],
            ["Cibulskis", "Černiauskienė", "Dagilė", "Žukauskas"],
        )


class Dublikatai(TestCase):
    def setUp(self):
        self.diena = date.today() + timedelta(days=7)
        self.kab = Cabinet.objects.create(pavadinimas="101")
        self.fizika = Subject.objects.create(pavadinimas="Fizika")
        self.chemija = Subject.objects.create(pavadinimas="Chemija")
        self.k6a = Klase.objects.create(pavadinimas="6A")
        self.k7b = Klase.objects.create(pavadinimas="7B")
        self.tevai = User.objects.create_user("tevai", "t@test.lt")

        # paliekamas - turi el. pasta ir yra seniausias
        self.ona = Teacher.objects.create(vardas="Ona", pavarde="Onaitė", email="ona@herojus.lt")
        self.ona.dalykai.add(self.fizika)
        # dublikatas su tuo paciu el. pastu ir savo duomenimis
        self.ona2 = Teacher.objects.create(vardas="Ona", pavarde="Onaite", email="ONA@herojus.lt", kabinetas=self.kab)
        self.ona2.dalykai.add(self.chemija)
        self.ona2.klases.add(self.k7b)
        # dublikatas be el. pasto, tuo paciu vardu
        self.ona3 = Teacher.objects.create(vardas="Ona", pavarde="Onaitė")

        self.k6a.aukletojas = self.ona2
        self.k6a.save()
        self.wh = WorkingHours.objects.create(date=self.diena, start_time=time(15, 0), end_time=time(16, 0), interval=15)
        self.wh.mokytojai.add(self.ona2)
        Reservation.objects.create(user=self.tevai, teacher=self.ona2, date=self.diena, time=time(15, 0))
        Reservation.objects.create(user=self.tevai, teacher=self.ona3, date=self.diena, time=time(15, 15))
        vadovas = Vadovas.objects.create(vardas="Jonas", pavarde="Vadovas", email="v@test.lt")
        VadovoRezervacija.objects.create(vadovas=vadovas, mokytojas=self.ona3, date=self.diena, time=time(14, 0))

    def test_visi_rysiai_aptarti(self):
        # atsiradus naujam rysiui su mokytoju ar klase, ji reikia perkelti ir sujungiant dublikatus
        self.assertEqual(
            {f.name for f in Teacher._meta.related_objects},
            {"auklejamos_klases", "working_hours", "reservation", "vadovo_rezervacijos"},
        )
        self.assertEqual({f.name for f in Klase._meta.related_objects}, {"teacher", "child"})

    def test_mokytoju_dublikatai_sujungiami_niekas_nedingsta(self):
        rezervaciju = Reservation.objects.count()
        planas = perziureti_importa([])
        self.assertEqual(len(planas["dublikatai"]), 2)
        # perziura nieko nekeicia
        self.assertEqual(Teacher.objects.count(), 3)

        rezultatas = vykdyti_importa([], planas["parasas"])
        self.assertTrue(rezultatas["ivykdyta"])
        self.assertEqual(list(Teacher.objects.all()), [self.ona])
        self.ona.refresh_from_db()
        self.assertEqual(self.ona.kabinetas, self.kab)
        self.assertEqual(sorted(d.pavadinimas for d in self.ona.dalykai.all()), ["Chemija", "Fizika"])
        self.assertEqual(list(self.ona.klases.all()), [self.k7b])
        self.assertEqual(Reservation.objects.count(), rezervaciju)
        self.assertEqual(Reservation.objects.filter(teacher=self.ona).count(), rezervaciju)
        self.assertEqual(VadovoRezervacija.objects.get().mokytojas, self.ona)
        self.assertEqual(list(self.wh.mokytojai.all()), [self.ona])
        self.k6a.refresh_from_db()
        self.assertEqual(self.k6a.aukletojas, self.ona)

    def test_rezervacija_tuo_paciu_laiku_neleidzia_sujungti(self):
        Reservation.objects.create(user=self.tevai, teacher=self.ona, date=self.diena, time=time(15, 0))
        busena = db_busena()
        planas = perziureti_importa([])
        self.assertIn("sujungti negalima", planas["klaidos"][0])
        rezultatas = vykdyti_importa([], planas["parasas"])
        self.assertFalse(rezultatas["ivykdyta"])
        self.assertEqual(db_busena(), busena)
        self.assertEqual(Teacher.objects.count(), 3)

    def test_bendravardziai_su_skirtingais_el_pastais_nesujungiami(self):
        Teacher.objects.filter(pk__in=[self.ona2.pk, self.ona3.pk]).delete()
        kita = Teacher.objects.create(vardas="Ona", pavarde="Onaitė", email="ona@gmail.com")
        planas = perziureti_importa([])
        self.assertEqual(planas["dublikatai"], [])
        self.assertIn("ona@gmail.com", planas["ispejimai"][0])
        vykdyti_importa([], planas["parasas"])
        self.assertTrue(Teacher.objects.filter(pk=kita.pk).exists())

    def test_klasiu_dublikatai(self):
        Teacher.objects.filter(pk__in=[self.ona2.pk, self.ona3.pk]).delete()
        k6a_2 = Klase.objects.create(pavadinimas="6 a", aukletojas=self.ona)
        vaikas = Child.objects.create(user=self.tevai, first_name="Jonas", last_name="Jonaitis", klase=k6a_2)
        self.ona.klases.add(k6a_2)

        planas = perziureti_importa([])
        self.assertEqual([(d["palikta"], d["istrinta"]) for d in planas["dublikatai"]], [("6A", "6 a")])
        vykdyti_importa([], planas["parasas"])

        self.assertFalse(Klase.objects.filter(pk=k6a_2.pk).exists())
        vaikas.refresh_from_db()
        self.assertEqual(vaikas.klase, self.k6a)
        self.assertEqual(list(self.ona.klases.all()), [self.k6a])
        self.k6a.refresh_from_db()
        # 6A auklėtojo neturejo - paima dublikato
        self.assertEqual(self.k6a.aukletojas, self.ona)

    def test_importas_nekuria_klases_dublikato(self):
        Teacher.objects.filter(pk__in=[self.ona2.pk, self.ona3.pk]).delete()
        importuoti([[1, "Ona", "Onaitė", "ona@herojus.lt", "", "", "6 a"]])
        # "6 a" atpazinta kaip esama 6A - naujos klases nesukurta
        self.assertEqual(sorted(Klase.objects.values_list("pavadinimas", flat=True)), ["6A", "7B"])
        self.k6a.refresh_from_db()
        self.assertEqual(self.k6a.aukletojas, self.ona)


class KabinetuSarasas(TestCase):
    def test_mokytojo_ir_vadovo_kabinetai_pagal_abecele(self):
        for pav in ["Windsor klasė", "Augustinos klasė", "Šviesa", "Isle of Skye klasė", "HAU klasė 3a.", "321", "101", "Cobra", "Čiurlionio"]:
            Cabinet.objects.create(pavadinimas=pav)
        tvarka = ["101", "321", "Augustinos klasė", "Cobra", "Čiurlionio", "HAU klasė 3a.", "Isle of Skye klasė", "Šviesa", "Windsor klasė"]

        Teacher.objects.create(vardas="Ona", pavarde="Onaitė", email="ona@test.lt")
        self.client.force_login(User.objects.create_user("ona", email="ona@test.lt"))
        laukas = self.client.get(reverse("teacher_dashboard")).context["wh_form"]["cabinet"]
        pasirinkimai = [str(p.choice_label) for p in laukas]
        # pirmas - "mokytojo numatytas kabinetas" (nebutina rinktis)
        self.assertEqual(pasirinkimai[0], "-- mokytojo numatytas kabinetas --")
        self.assertEqual(pasirinkimai[1:], tvarka)

        Vadovas.objects.create(vardas="Jonas", pavarde="Jonaitis", email="jonas@test.lt")
        self.client.force_login(User.objects.create_user("jonas", email="jonas@test.lt"))
        laukas = self.client.get(reverse("vadovas_dashboard")).context["wh_form"]["cabinet"]
        self.assertEqual([str(p.choice_label) for p in laukas][1:], tvarka)


class MokytojoGrafikoKlases(TestCase):
    def test_virsuje_rodomos_klases(self):
        mokytojas = Teacher.objects.create(vardas="Ona", pavarde="Onaitė", email="ona@test.lt")
        self.client.force_login(User.objects.create_user("ona", email="ona@test.lt"))
        r = self.client.get(reverse("teacher_dashboard"))
        self.assertContains(r, "nepriskirtos")

        k10 = Klase.objects.create(pavadinimas="10S")
        k6 = Klase.objects.create(pavadinimas="6F", aukletojas=mokytojas)
        mokytojas.klases.add(k10, k6)
        r = self.client.get(reverse("teacher_dashboard"))
        self.assertEqual([k.pavadinimas for k in r.context["auklejamos_klases"]], ["6F"])
        self.assertEqual([k.pavadinimas for k in r.context["mokomos_klases"]], ["6F", "10S"])
        self.assertNotContains(r, "nepriskirtos")
